"""Builds the submission deck: a 1:1 PowerPoint copy of the five business slides in demo/index.html.

    python pitch/build_deck.py            # writes pitch/ITSG-Spec-Pipeline.pptx
    python pitch/build_deck.py --pdf      # also exports the PDF through PowerPoint (Windows), else LibreOffice

Text, notes and colours come from the SLIDES array and the light-theme :root tokens of demo/index.html.
"""
import os
import shutil
import subprocess
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
IMG = os.path.join(ROOT, "images")
OUT = os.path.join(HERE, "ITSG-Spec-Pipeline.pptx")

NAME = "§pecs"
TAGLINE = "From specification to certified test, in one session."

# light theme of demo/index.html
BG, SURFACE, LINE = "F6F4F4", "FFFFFF", "E0D9D9"
FG, FG2, MUTED = "1C1819", "3D3537", "6D6366"
ACCENT, ACCENT_INK = "B0142F", "FFFFFF"
DISPLAY, BODY, EMOJI = "Arial Narrow", "Segoe UI", "Segoe UI Emoji"

W, H = 13.333, 7.5


def rgb(hex_):
    return RGBColor.from_string(hex_)


def text(slide, x, y, w, h, runs, size, color=FG, font=BODY, bold=False, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, spacing=None, line=None):
    """A text box. `runs` is a string, or a list of paragraphs, each a list of (text, {overrides})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paras = [[(runs, {})]] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        for t, o in para:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.name = o.get("font", font)
            f.color.rgb = rgb(o.get("color", color))
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
    return box


def card(slide, x, y, w, h, fill=SURFACE, line=LINE, radius=0.12):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def soft_shadow(shape):
    """The app's --shadow: a soft, low-opacity drop shadow."""
    sp = shape._element.spPr
    eff = sp.makeelement(qn("a:effectLst"), {})
    sh = eff.makeelement(qn("a:outerShdw"), {"blurRad": "228600", "dist": "57150", "dir": "5400000", "algn": "t", "rotWithShape": "0"})
    clr = sh.makeelement(qn("a:srgbClr"), {"val": "281418"})
    clr.append(clr.makeelement(qn("a:alpha"), {"val": "16000"}))
    sh.append(clr)
    eff.append(sh)
    sp.append(eff)


def photo(slide, path, x, y, w, h, alt):
    """object-fit: cover with rounded corners (border-radius 18px)."""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    box, img = w / h, iw / ih
    if img > box:  # too wide: crop the sides
        c = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = c
    elif img < box:  # too tall: crop top and bottom
        c = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = c
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    av = pic._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val 6000"}))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    soft_shadow(pic)
    return pic


def new_slide(prs, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    s.notes_slide.notes_text_frame.text = "\n".join(notes)
    return s


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = NAME
    prs.core_properties.subject = TAGLINE

    # 1 · Welcome
    s = new_slide(prs, ["We built §pecs for the ITSG challenge.",
                        "In one line: from the yearly specification to certified test cases."])
    lw = 1.45 * 681 / 293
    logo = s.shapes.add_picture(os.path.join(IMG, "ITSG-logo.png"), Inches((W - lw) / 2), Inches(1.75), height=Inches(1.45))
    logo._element.nvPicPr.cNvPr.set("descr", "ITSG")
    text(s, 0.8, 3.5, W - 1.6, 1.25, NAME, 66, font=DISPLAY, bold=True)
    text(s, 1.5, 4.85, W - 3, 0.6, TAGLINE, 24, color=FG2)

    # 2 · The hustle
    s = new_slide(prs, ["Every December a new specification lands: the Verfahrensbeschreibung alone is 168 pages, "
                        "around 400 with the other EEL documents.",
                        "Experts read every line, highlight it by hand and rewrite the test cases."])
    photo(s, os.path.join(IMG, "tired-patrick-meme.jpeg"), 1.0, 1.1, 5.3, 5.3, "Patrick Star with red, tired eyes")
    text(s, 7.0, 1.55, 5.5, 2.6, [[("400 pages.", {})], [("Every year.", {})], [("By hand.", {})]], 50,
         font=DISPLAY, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.BOTTOM, line=1.0)
    text(s, 7.0, 4.45, 5.0, 1.6, [[("A new specification lands every December. ", {}), ("~400 pages.", {"bold": True, "color": FG}),
                                   (" Someone reads every line, and writes every test case.", {})]], 20,
         color=FG2, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.25)

    # 3 · Where the hours go
    s = new_slide(prs, ["Four manual steps: read, extract the rules, write 3,000–4,000 test cases, and every vendor's "
                        "developer types them in again.",
                        "Zero error tolerance: one wrong expected value and a vendor fails certification."])
    text(s, 0.8, 1.15, W - 1.6, 0.35, "WHERE THE HOURS GO", 14, color=ACCENT, bold=True, spacing=200)
    text(s, 0.8, 1.6, W - 1.6, 0.9, "Four steps, all by hand", 44, font=DISPLAY, bold=True)
    steps = [("📖", "Read the spec", "~400 pages"), ("✍️", "Extract rules", "by hand, § by §"),
             ("☑️", "Write test cases", "3,000–4,000 cases"), ("⌨️", "Dev types them in", "manually, per vendor")]
    cw, gap, x0, cy, ch = 2.75, 0.27, 0.8, 2.85, 2.2
    for i, (ic, title, sub) in enumerate(steps):
        x = x0 + i * (cw + gap)
        card(s, x, cy, cw, ch)
        text(s, x + 0.3, cy + 0.3, cw - 0.6, 0.6, ic, 30, font=EMOJI, align=PP_ALIGN.LEFT)
        text(s, x + 0.3, cy + 1.0, cw - 0.6, 0.45, title, 22, font=DISPLAY, bold=True, align=PP_ALIGN.LEFT)
        text(s, x + 0.3, cy + 1.45, cw - 0.6, 0.4, sub, 15, color=ACCENT, bold=True, align=PP_ALIGN.LEFT)
        if i < len(steps) - 1:
            text(s, x + cw - 0.1, cy + ch / 2 - 0.25, gap + 0.2, 0.5, "→", 24, color=ACCENT, bold=True)
    bar = card(s, x0, 5.45, 4 * cw + 3 * gap, 0.9, fill=ACCENT, line=None, radius=0.14)
    text(s, x0, 5.45, 4 * cw + 3 * gap, 0.9, "Every year. By hand. Zero error tolerance.", 28, color=ACCENT_INK,
         font=DISPLAY, bold=True)

    # 4 · Our solution (the notes also carry data, model and logic for the submission)
    s = new_slide(prs, ["§pecs: the AI reads and drafts, you approve, tested code calculates.",
                        "Every value traces back to a sentence in the spec. Let me show you.",
                        "Data: the government's v13.4 Verfahrensbeschreibung PDF (all 220 sections), last year's v13.3 "
                        "from the archive, and ITSG's own reference case and TEEL0-Test1-DBVO file.",
                        "Model: vllm/Kolibri-1 (Aleph Alpha) through the hessian.AI gateway, with tool calls. It only "
                        "extracts parameters together with verbatim spec quotes; it never computes an expected value.",
                        "Logic: structural § retrieval (headings, TOC, internal links) → deterministic diff and change "
                        "log → AI extraction → 5 guards (quote verbatim, formula parses, oracle = AI value, number in "
                        "quote, formula was run) → examiner approves → tested functions compute every expected value."])
    text(s, 0.8, 1.3, W - 1.6, 0.35, "OUR SOLUTION", 14, color=ACCENT, bold=True, spacing=200)
    text(s, 0.8, 1.8, W - 1.6, 1.2, NAME, 66, font=DISPLAY, bold=True)
    text(s, 1.5, 3.05, W - 3, 0.6, TAGLINE, 24, color=FG2)
    promises = [("🤖", "AI reads the spec"), ("👤", "You approve"), ("🔗", "Every value traces to a §")]
    pw, pg = 3.5, 0.3
    px0 = (W - (3 * pw + 2 * pg)) / 2
    for i, (ic, label) in enumerate(promises):
        x = px0 + i * (pw + pg)
        card(s, x, 4.15, pw, 1.75)
        text(s, x, 4.35, pw, 0.65, ic, 32, font=EMOJI)
        text(s, x + 0.2, 5.05, pw - 0.4, 0.5, label, 19, bold=True)

    # 5 · Before → after · Thank you
    s = new_slide(prs, ["Before: weeks of reading and typing. After: one session, every value traced to the spec.",
                        "Thank you. Questions?"])
    iw, ih, iy = 5.0, 3.3, 0.7
    photo(s, os.path.join(IMG, "tired-guy-meme.jpeg"), 0.9, iy, iw, ih, "A tired man at a desk late at night")
    photo(s, os.path.join(IMG, "happy-squirrel-meme.jpeg"), W - 0.9 - iw, iy, iw, ih,
          "A happy squirrel riding a unicorn under a rainbow")
    text(s, 0.9 + iw, iy + ih / 2 - 0.5, W - 1.8 - 2 * iw, 1.0, "→", 56, color=ACCENT, bold=True)
    for x, head, sub in ((0.9, "Before", "Weeks of reading. Thousands of test cases, written and typed in by hand."),
                         (W - 0.9 - iw, "After", "Spec in, certified tests out: in one session, every value traced to the spec.")):
        text(s, x, iy + ih + 0.15, iw, 0.45, head, 24, font=DISPLAY, bold=True, align=PP_ALIGN.LEFT)
        text(s, x, iy + ih + 0.6, iw, 0.7, sub, 15, color=MUTED, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    text(s, 0.8, 5.7, W - 1.6, 1.1, "Thank you!", 60, font=DISPLAY, bold=True)

    prs.save(OUT)
    print("wrote", OUT)


def export_pdf():
    pdf = OUT[:-5] + ".pdf"
    if sys.platform == "win32":
        ps = (f"$p = New-Object -ComObject PowerPoint.Application; "
              f"$d = $p.Presentations.Open('{OUT}', $true, $false, $false); "
              f"$d.SaveAs('{pdf}', 32); $d.Close(); $p.Quit()")
        if subprocess.run(["powershell", "-NoProfile", "-Command", ps]).returncode == 0 and os.path.exists(pdf):
            print("wrote", pdf, "(PowerPoint)")
            return
    soffice = shutil.which("soffice") or r"C:\Program Files\LibreOffice\program\soffice.exe"
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", HERE, OUT], check=True)
    print("wrote", pdf, "(LibreOffice)")


if __name__ == "__main__":
    build()
    if "--pdf" in sys.argv:
        export_pdf()
